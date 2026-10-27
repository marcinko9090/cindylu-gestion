import unittest,tempfile,os,sys,json,datetime as dt,concurrent.futures
sys.path.insert(0,str(__import__('pathlib').Path(__file__).resolve().parents[1]))
tmp=tempfile.TemporaryDirectory();os.environ['CINDY_DATA']=tmp.name
import app as a

class FinanceTests(unittest.TestCase):
    def setUp(self):
        with a.db(True) as c:
            for t in ['statement_sends','payment_notices','allocations','payments','installments','messages','receipts','orders','clients','users']:c.execute('DELETE FROM '+t)
            self.cid=c.execute("INSERT INTO clients(name,phone) VALUES('Javiera González','56900000001')").lastrowid
            self.oid=c.execute("INSERT INTO orders(client_id,number,purchased,total,items,updated) VALUES(?,'#1','2026-10-01',100001,'[]',?)",(self.cid,a.now())).lastrowid
        self.http=a.app.test_client()
    def plan(self,n=3,first='2027-01-31'):
        with a.db(True) as c:a.make_plan(c,{'order_id':self.oid,'count':n,'first':first})
    def payment(self,n,operation='abcdefghijklmnop',receipt=None):
        with a.db(True) as c:return a.pay(c,{'order_id':self.oid,'amount':n,'operation':operation,'receipt_id':receipt})
    def test_statement_filters_and_reconciles(self):
        from statements import snapshot,pdf
        self.plan();self.payment(40000)
        with a.db(True) as c:
            c.execute("INSERT INTO orders(client_id,number,purchased,total,base_paid,items,updated) VALUES(?,'PAID','2026-01-01',100,100,'[]',?)",(self.cid,a.now()))
            c.execute("INSERT INTO orders(client_id,number,purchased,total,items,status,updated) VALUES(?,'CANCELLED','2026-01-01',100,'[]','cancelled',?)",(self.cid,a.now()))
        with a.db() as c:d=snapshot(c,self.cid,a.balance,'26/10/2026 12:00')
        self.assertEqual(len(d['orders']),1);self.assertEqual(d['balance'],60001)
        self.assertEqual(d['total']-d['paid'],d['balance'])
        self.assertEqual(sum(i['amount']-i['paid'] for i in d['orders'][0]['installments']),60001)
        self.assertTrue(pdf(d).getvalue().startswith(b'%PDF'))
        self.assertEqual(self.http.get('/api/clients/1/statement.pdf').status_code,401)
        with self.http.session_transaction() as sess:sess['user']='test'
        self.assertEqual(self.http.get('/api/clients/'+str(self.cid)+'/statement.pdf').mimetype,'application/pdf')
        with a.db(True) as c:c.execute('UPDATE payments SET reversed=1')
        with a.db() as c:self.assertEqual(snapshot(c,self.cid,a.balance,'now')['balance'],100001)
    def test_statement_empty(self):
        from statements import snapshot,pdf
        with a.db(True) as c:c.execute('UPDATE orders SET base_paid=total')
        with a.db() as c:d=snapshot(c,self.cid,a.balance,'now')
        self.assertEqual(d['orders'],[]);self.assertEqual(d['balance'],0)
        self.assertTrue(pdf(d).getvalue().startswith(b'%PDF'))
    def statement_login(self):
        with self.http.session_transaction() as session:session['user']=1;session['csrf']='test'
        with a.db(True) as c:
            c.execute('UPDATE clients SET consent=1 WHERE id=?',(self.cid,))
            a.setval(c,'sending',True);a.setval(c,'statement_template','estado_cuenta');a.setval(c,'statement_language','es_CL')
    def statement_post(self,preview,operation='statement-test-000001'):
        return self.http.post('/api/statement-sends',json={'client_id':self.cid,'digest':preview['digest'],'operation':operation},headers={'X-CSRF-Token':'test'})
    def test_statement_demo_no_network_and_private_copy(self):
        from unittest.mock import patch
        self.assertEqual(self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').status_code,401)
        self.statement_login()
        with patch.object(a,'DEMO',True),patch.object(a,'graph') as graph,patch('wa_documents.upload_pdf') as upload:
            preview=self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').json
            r=self.statement_post(preview);self.assertEqual(r.status_code,200);self.assertEqual(r.json['status'],'demo')
            self.assertEqual(self.statement_post(preview).json['id'],r.json['id']);graph.assert_not_called();upload.assert_not_called()
        file=self.http.get('/api/statement-sends/'+str(r.json['id'])+'/file');self.assertTrue(file.data.startswith(b'%PDF'));file.close()
        self.http.post('/api/logout',headers={'X-CSRF-Token':'test'})
        self.assertEqual(self.http.get('/api/statement-sends/'+str(r.json['id'])+'/file').status_code,401)
        with a.db() as c:self.assertEqual(a.balance(c,self.oid),100001)
    def test_statement_real_document_idempotent(self):
        from unittest.mock import patch
        self.statement_login()
        with a.db(True) as c:c.execute("INSERT INTO messages(client_id,direction,body,wa_id,status,created) VALUES(?,'in','Saldo por favor','in-1','received',?)",(self.cid,a.now()))
        with patch.dict(os.environ,{'WA_ACCESS_TOKEN':'fake','WA_PHONE_NUMBER_ID':'123','WA_APP_SECRET':'fake'}),patch('wa_documents.upload_pdf',return_value='media-1') as upload,patch.object(a,'graph',return_value={'messages':[{'id':'out-1'}]}) as graph:
            preview=self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').json
            self.assertEqual(preview['mode'],'document');r=self.statement_post(preview);self.assertEqual(r.status_code,200)
            self.assertTrue(self.statement_post(preview).json['duplicate']);self.assertEqual(upload.call_count,1);self.assertEqual(graph.call_count,1)
            sent=graph.call_args.args[1];self.assertEqual(sent['type'],'document');self.assertEqual(sent['to'],'56900000001');self.assertEqual(sent['document']['id'],'media-1');self.assertNotIn('link',sent['document'])
    def test_statement_template_and_stale_balance(self):
        from unittest.mock import patch
        self.statement_login();self.plan()
        with patch.dict(os.environ,{'WA_ACCESS_TOKEN':'fake','WA_PHONE_NUMBER_ID':'123','WA_APP_SECRET':'fake'}),patch('wa_documents.upload_pdf',return_value='media-1') as upload,patch.object(a,'graph',return_value={'messages':[{'id':'out-template'}]}) as graph:
            preview=self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').json
            self.payment(1000)
            self.assertEqual(self.statement_post(preview).status_code,400);upload.assert_not_called()
            preview=self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').json
            self.assertEqual(self.statement_post(preview).status_code,200)
            payload=graph.call_args.args[1];self.assertEqual(payload['type'],'template')
            components=payload['template']['components'];self.assertEqual(components[0]['parameters'][0]['type'],'document');self.assertEqual(components[1]['parameters'][1]['text'],'$99.001')
    def test_statement_blockers_before_upload(self):
        from unittest.mock import patch
        self.statement_login()
        with patch.dict(os.environ,{'WA_ACCESS_TOKEN':'fake','WA_PHONE_NUMBER_ID':'123','WA_APP_SECRET':'fake'}),patch('wa_documents.upload_pdf') as upload:
            for change in ["UPDATE clients SET consent=0","UPDATE clients SET consent=1,phone=''","UPDATE clients SET phone='56900000001'",]:
                with a.db(True) as c:
                    c.execute(change)
                    if change.endswith("'56900000001'"):a.setval(c,'statement_template','')
                p=self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').json
                self.assertTrue(p['blocked']);self.assertEqual(self.statement_post(p).status_code,400)
            upload.assert_not_called()
    def test_statement_ambiguous_failure_not_retried(self):
        from unittest.mock import patch
        self.statement_login()
        with patch.dict(os.environ,{'WA_ACCESS_TOKEN':'fake','WA_PHONE_NUMBER_ID':'123','WA_APP_SECRET':'fake'}),patch('wa_documents.upload_pdf',return_value='media-1'),patch.object(a,'graph',side_effect=TimeoutError) as graph:
            p=self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').json
            self.assertEqual(self.statement_post(p).status_code,502)
            self.assertEqual(self.statement_post(p).json['status'],'review');self.assertEqual(graph.call_count,1)
    def test_statement_media_multipart(self):
        from unittest.mock import patch,MagicMock
        from wa_documents import upload_pdf
        response=MagicMock();response.__enter__.return_value.read.return_value=b'{"id":"media-test"}'
        with patch.dict(os.environ,{'CINDY_DEMO':'0','WA_ACCESS_TOKEN':'private-test','WA_PHONE_NUMBER_ID':'123'}),patch('urllib.request.urlopen',return_value=response) as urlopen:
            self.assertEqual(upload_pdf(b'%PDF-test','test.pdf'),'media-test')
            req=urlopen.call_args.args[0];self.assertIn(b'name="messaging_product"',req.data);self.assertIn(b'application/pdf',req.data);self.assertIn(b'%PDF-test',req.data)
        with patch.dict(os.environ,{'CINDY_DEMO':'1'}),patch('urllib.request.urlopen') as urlopen:
            with self.assertRaises(ValueError):upload_pdf(b'%PDF-test','test.pdf')
            urlopen.assert_not_called()
    def test_statement_webhook_delivery_does_not_regress(self):
        from unittest.mock import patch
        self.statement_login()
        with patch.dict(os.environ,{'WA_ACCESS_TOKEN':'fake','WA_PHONE_NUMBER_ID':'123','WA_APP_SECRET':'test-secret'}),patch('wa_documents.upload_pdf',return_value='media-1'),patch.object(a,'graph',return_value={'messages':[{'id':'out-status'}]}):
            preview=self.http.get('/api/clients/'+str(self.cid)+'/statement-preview').json
            r=self.statement_post(preview);self.assertEqual(r.status_code,200)
            for status in ['read','sent','delivered']:
                payload=json.dumps({'entry':[{'changes':[{'value':{'metadata':{'phone_number_id':'123'},'statuses':[{'id':'out-status','status':status}]}}]}]}).encode()
                signature='sha256='+a.hmac.new(b'test-secret',payload,a.hashlib.sha256).hexdigest()
                self.assertEqual(self.http.post('/webhooks/whatsapp',data=payload,content_type='application/json',headers={'X-Hub-Signature-256':signature}).status_code,200)
            with a.db() as c:self.assertEqual(c.execute('SELECT status FROM statement_sends WHERE id=?',(r.json['id'],)).fetchone()[0],'read')
    def test_rounding_and_month_end(self):
        self.plan()
        with a.db() as c:
            ins=a.rows(c,'SELECT * FROM installments');self.assertEqual([i['due'] for i in ins],['2027-01-31','2027-02-28','2027-03-31']);self.assertEqual(sum(i['amount'] for i in ins),100001)
    def test_partial_covers_multiple_installments_once(self):
        self.plan();self.payment(40000);self.payment(40000)
        with a.db() as c:
            self.assertEqual(a.balance(c,self.oid),60001);self.assertEqual(c.execute('SELECT SUM(amount) FROM allocations').fetchone()[0],40000);self.assertEqual(c.execute('SELECT COUNT(*) FROM allocations').fetchone()[0],2)
    def test_concurrent_overpayment(self):
        self.plan()
        def pay(i):
            try:self.payment(70000,'concurrent000000'+str(i));return True
            except ValueError:return False
        with concurrent.futures.ThreadPoolExecutor(2) as pool:results=list(pool.map(pay,[1,2]))
        self.assertEqual(sum(results),1)
    def test_receipt_duplicate_and_reversal(self):
        self.plan()
        with a.db(True) as c:rid=c.execute("INSERT INTO receipts(client_id,filename,mime,created) VALUES(?,'x','image/png',?)",(self.cid,a.now())).lastrowid
        pid=self.payment(20000,receipt=rid)
        with self.assertRaises(ValueError):self.payment(20000,'anotheroperation00',rid)
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        r=self.http.post('/api/payments/'+str(pid)+'/reverse',json={'reason':'Monto equivocado'},headers={'X-CSRF-Token':'test'});self.assertEqual(r.status_code,200)
        self.payment(10000,'reconfirmoperation',rid)
        with a.db() as c:self.assertEqual(a.balance(c,self.oid),90001)
    def test_login_setup_csrf(self):
        self.assertEqual(self.http.get('/api/state').status_code,401)
        self.assertEqual(self.http.post('/api/setup',json={'email':'test@example.com','password':'long-password-123'}).status_code,200)
        self.assertEqual(self.http.post('/api/setup',json={'email':'test@example.com','password':'long-password-123'}).status_code,400)
        r=self.http.post('/api/login',json={'email':'test@example.com','password':'long-password-123'});self.assertEqual(r.status_code,200)
        self.assertEqual(self.http.post('/api/plan',json={}).status_code,403)
        self.assertEqual(self.http.get('/api/state').status_code,200)
    def test_shopify_idempotence_and_review(self):
        order={'id':'gid://shopify/Order/1','name':'#100','createdAt':a.now(),'cancelledAt':None,'customer':{'id':'gid://shopify/Customer/1','displayName':'Javiera','phone':'56900000001'},'currentTotalPriceSet':{'shopMoney':{'amount':'150000','currencyCode':'CLP'}},'totalOutstandingSet':{'shopMoney':{'amount':'120000','currencyCode':'CLP'}},'lineItems':{'nodes':[{'title':'Labial','quantity':1}]}}
        with a.db(True) as c:a.import_shop_order(c,order);a.import_shop_order(c,order)
        with a.db() as c:
            os_=a.rows(c,"SELECT * FROM orders WHERE source='shopify'");self.assertEqual(len(os_),1);self.assertEqual(os_[0]['base_paid'],30000);self.assertEqual(os_[0]['client_id'],self.cid)
        order['totalOutstandingSet']['shopMoney']['amount']='100000'
        with a.db(True) as c:a.import_shop_order(c,order)
        with a.db() as c:
            o=a.one(c,"SELECT * FROM orders WHERE source='shopify'");self.assertTrue(o['review']);self.assertEqual(o['base_paid'],30000)
    def test_bad_money_and_payment_without_installments(self):
        for n in [-1,1.5,'NaN','Infinity']:
            with self.assertRaises(ValueError):a.amount(n)
        self.payment(100)
        with a.db() as c:self.assertEqual(a.balance(c,self.oid),99901)
    def test_signature_rejected(self):
        self.assertEqual(self.http.post('/webhooks/whatsapp',json={}).status_code,403)
    def test_upload_validation(self):
        with self.assertRaises(ValueError):a.save_media(b'<script>alert(1)</script>')
    def test_plan_cannot_overwrite_payments(self):
        self.plan();self.payment(100)
        with self.assertRaises(ValueError):self.plan(4)

    def test_excel_preview_and_duplicate_import(self):
        import io
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        r=self.http.post('/api/import/preview',data={'file':(io.BytesIO('nombre;telefono;deuda\nMaría;56911111111;20000\n'.encode()),'clientes.csv')},headers={'X-CSRF-Token':'test'})
        self.assertEqual(r.status_code,200);data=[x['data'] for x in r.json['rows']]
        for expected in [1,0]:
            r=self.http.post('/api/import/confirm',json={'rows':data},headers={'X-CSRF-Token':'test'});self.assertEqual(r.json['inserted'],expected)
        with a.db() as c:self.assertEqual(c.execute("SELECT COUNT(*) FROM orders WHERE source='excel'").fetchone()[0],1)
    def test_signed_webhook_receipt_idempotency(self):
        import hashlib,hmac
        os.environ['WA_APP_SECRET']='test-signature-secret';os.environ['WA_PHONE_NUMBER_ID']='testnumber'
        payload={'entry':[{'changes':[{'value':{'metadata':{'phone_number_id':'testnumber'},'messages':[{'id':'wamid.test','from':'56900000001','type':'image','timestamp':str(int(__import__('time').time())),'image':{'id':'testmedia','mime_type':'image/png'}}]}}]}]}
        raw=json.dumps(payload).encode();sig='sha256='+hmac.new(b'test-signature-secret',raw,hashlib.sha256).hexdigest()
        for _ in range(2):
            r=self.http.post('/webhooks/whatsapp',data=raw,headers={'Content-Type':'application/json','X-Hub-Signature-256':sig});self.assertEqual(r.status_code,200)
        with a.db() as c:
            self.assertEqual(c.execute('SELECT COUNT(*) FROM receipts').fetchone()[0],1);self.assertEqual(a.balance(c,self.oid),100001)
        os.environ.pop('WA_APP_SECRET');os.environ.pop('WA_PHONE_NUMBER_ID')
    def test_xlsx_numeric_phone(self):
        import io
        from openpyxl import Workbook
        book=Workbook();book.active.append(['nombre','telefono','deuda','fecha']);book.active.append(['Ana',56922222222,10000,dt.date(2026,10,26)]);buf=io.BytesIO();book.save(buf);buf.seek(0)
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        r=self.http.post('/api/import/preview',data={'file':(buf,'clientes.xlsx')},headers={'X-CSRF-Token':'test'})
        self.assertEqual(r.status_code,200);self.assertEqual(r.json['rows'][0]['data']['phone'],'56922222222');self.assertEqual(r.json['rows'][0]['data']['date'],'2026-10-26')
    def test_cross_origin_and_private_media(self):
        self.assertEqual(self.http.get('/api/receipts/1/file').status_code,401)
        self.assertEqual(self.http.post('/api/setup',json={},headers={'Origin':'https://evil.example'}).status_code,403)

    def test_payment_notice_is_atomic_and_not_duplicated(self):
        self.plan()
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        payload={'order_id':self.oid,'amount':10000,'operation':'notificationoperation','notify':True}
        for _ in range(2):
            r=self.http.post('/api/payments',json=payload,headers={'X-CSRF-Token':'test'});self.assertEqual(r.status_code,200)
        with a.db() as c:
            notices=a.rows(c,'SELECT * FROM payment_notices');self.assertEqual(len(notices),1);self.assertIn('$90.001',notices[0]['body']);self.assertIn('Javiera',notices[0]['body']);self.assertEqual(a.balance(c,self.oid),90001)
        self.http.post('/api/payments/'+str(r.json['id'])+'/reverse',json={'reason':'Corrección del monto'},headers={'X-CSRF-Token':'test'})
        with a.db() as c:self.assertEqual(c.execute('SELECT status FROM payment_notices').fetchone()[0],'cancelled')
    def test_failed_notice_does_not_undo_or_resend_payment(self):
        from unittest.mock import patch
        self.plan()
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        self.http.post('/api/payments',json={'order_id':self.oid,'amount':10000,'operation':'failednoticeoperation','notify':True},headers={'X-CSRF-Token':'test'})
        with a.db(True) as c:a.setval(c,'sending',True)
        with patch.dict(os.environ,{'WA_ACCESS_TOKEN':'test-only','WA_PHONE_NUMBER_ID':'test-only'}),patch.object(a,'send_message',side_effect=ValueError('Prueba de fallo')) as send:
            a.process_payment_notices();a.process_payment_notices();self.assertEqual(send.call_count,1)
        with a.db(True) as c:
            self.assertEqual(a.balance(c,self.oid),90001);self.assertEqual(c.execute('SELECT status FROM payment_notices').fetchone()[0],'review');a.setval(c,'sending',False)
    def test_notice_off_creates_no_message(self):
        self.plan()
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        self.http.post('/api/payments',json={'order_id':self.oid,'amount':10000,'operation':'nonotificationoperation','notify':False},headers={'X-CSRF-Token':'test'})
        with a.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM payment_notices').fetchone()[0],0)

    def test_shopify_paid_order_imported_as_credit_and_adjustable(self):
        with a.db(True) as c:
            a.setval(c,'shopify_credit_mode','credit_when_paid')
            a.setval(c,'shopify_default_installments',3)
            order={'id':'gid://shopify/Order/999','name':'#999','createdAt':'2026-10-26T12:00:00Z','cancelledAt':None,'customer':{'id':'gid://shopify/Customer/99','displayName':'Carolina Silva','phone':'56988887777'},'currentTotalPriceSet':{'shopMoney':{'amount':'90000','currencyCode':'CLP'}},'totalOutstandingSet':{'shopMoney':{'amount':'0','currencyCode':'CLP'}},'lineItems':{'nodes':[{'title':'Vestido Gala','quantity':1}]}}
            a.import_shop_order(c,order)
            a.setval(c,'shopify_default_installments',0)
        with a.db() as c:
            o=a.one(c,"SELECT * FROM orders WHERE number='#999'")
            self.assertEqual(o['base_paid'],0)
            self.assertEqual(a.balance(c,o['id']),90000)
            ins=a.rows(c,'SELECT * FROM installments WHERE order_id=? ORDER BY number',(o['id'],))
            self.assertEqual(len(ins),0)
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        # Pay 30000 then adjust plan to 2 installments and verify payment is preserved and re-allocated
        with a.db(True) as c:a.pay(c,{'order_id':o['id'],'amount':30000,'operation':'shopcreditpay00001'})
        r=self.http.post('/api/orders/'+str(o['id'])+'/adjust',json={'total':90000,'base_paid':10000,'count':2,'first':'2026-11-26','frequency':'monthly'},headers={'X-CSRF-Token':'test'})
        self.assertEqual(r.status_code,200)
        with a.db() as c:
            self.assertEqual(a.balance(c,o['id']),50000)
            alloc_sum=c.execute('SELECT COALESCE(SUM(a.amount),0) FROM allocations a JOIN installments i ON i.id=a.installment_id WHERE i.order_id=?',(o['id'],)).fetchone()[0]
            self.assertEqual(alloc_sum,30000)

    def test_cindy_excel_multisheet_review_and_confirm(self):
        import io
        from openpyxl import Workbook
        wb=Workbook()
        ws1=wb.active;ws1.title='DATOS CLIENTAS'
        ws1.append(['COD','NOMBRE','TRABAJO','TELEFONO'])
        ws1.append([10,'Claudia Soto','Hospital','977776666'])
        ws2=wb.create_sheet('OCTUBRE')
        ws2.append(['COD','NOMBRE','SALDO VIGENTE'])
        ws2.append([10,'Claudia Soto',60000])
        buf=io.BytesIO();wb.save(buf);buf.seek(0)
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        rev=self.http.post('/api/import/cindy-review',data={'files':[(buf,'CXC_2026.xlsx')]},headers={'X-CSRF-Token':'test'})
        self.assertEqual(rev.status_code,200)
        imp=rev.json['importable_sources']
        self.assertEqual(len(imp),1)
        conf=self.http.post('/api/import/cindy-confirm',json={'records':imp[0]['records'],'source_label':'OCTUBRE','installments_count':3,'first_due':'2026-11-15','frequency':'monthly'},headers={'X-CSRF-Token':'test'})
        self.assertEqual(conf.status_code,200)
        self.assertEqual(conf.json['created_orders'],1)
        with a.db() as c:
            cl=a.one(c,"SELECT * FROM clients WHERE phone='56977776666'")
            self.assertEqual(cl['name'],'Claudia Soto')
            o=a.one(c,'SELECT * FROM orders WHERE client_id=?',(cl['id'],))
            self.assertEqual(a.balance(c,o['id']),60000)
            ins=a.rows(c,'SELECT * FROM installments WHERE order_id=?',(o['id'],))
            self.assertEqual(len(ins),0)

    def test_cindy_ledger_fifo_orders_and_payments_import(self):
        import io
        from openpyxl import Workbook
        wb=Workbook()
        ws1=wb.active;ws1.title='DATOS CLIENTAS'
        ws1.append(['oo','CLIENTA','LUGAR DE TRABAJO','CELULAR'])
        ws1.append([2543,'Alejandra Bravo Caceres','CLC',98765432])
        ws2=wb.create_sheet('CLUB CINDYLÚ')
        ws2.append(['COD','CLIENTA','FECHA','CODIGO','DETALLE','MONTO'])
        ws2.append([2543,'Alejandra Bravo Caceres','2026-07-30',6044,'PEDIDO',129980])
        ws2.append([2543,'Alejandra Bravo Caceres','2026-09-08','A','ABONO',-20000])
        ws2.append([2543,'Alejandra Bravo Caceres','2026-10-07','A','ABONO',-40000])
        ws3=wb.create_sheet('OCTUBRE')
        ws3.append(['COD','clienta','LUGAR DE TRABAJO','ULTIMO ABONO','FECHA ','SALDO VIGENTE','MONTO','FECHA','SALDO'])
        ws3.append([2543,'Alejandra Bravo Caceres','CLC',40000,'2026-10-07',69980,-10000,'2026-10-20',None])
        buf=io.BytesIO();wb.save(buf);buf.seek(0)
        with self.http.session_transaction() as s:s['user']=1;s['csrf']='test'
        rev=self.http.post('/api/import/cindy-review',data={'files':[(buf,'CLUB_CINDYLU.xlsx')]},headers={'X-CSRF-Token':'test'})
        self.assertEqual(rev.status_code,200)
        imp=next(x for x in rev.json['importable_sources'] if x['id']=='reconciled::cindy_master')
        self.assertEqual(len(imp['records']),1)
        self.assertEqual(imp['records'][0]['order_number'],'#6044')
        self.assertEqual(imp['records'][0]['balance'],59980)
        conf=self.http.post('/api/import/cindy-confirm',json={'records':imp['records'],'source_label':'Cartera Conciliada','installments_count':3,'first_due':'2026-08-30','frequency':'monthly'},headers={'X-CSRF-Token':'test'})
        self.assertEqual(conf.status_code,200)
        self.assertEqual(conf.json['created_payments'],3)
        with a.db() as c:
            cl=a.one(c,"SELECT * FROM clients WHERE phone='56998765432'")
            o=a.one(c,"SELECT * FROM orders WHERE client_id=? AND number='#6044'",(cl['id'],))
            self.assertEqual(a.balance(c,o['id']),59980)
            p_cnt=c.execute('SELECT COUNT(*) FROM payments WHERE order_id=? AND reversed=0',(o['id'],)).fetchone()[0]
            self.assertEqual(p_cnt,3)

if __name__=='__main__':unittest.main(verbosity=2)

